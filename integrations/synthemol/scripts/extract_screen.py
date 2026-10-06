"""Read the original XLSX; export observed growth, never model predictions."""
import argparse
import csv
import json
from pathlib import Path
from openpyxl import load_workbook


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source", type=Path)
    ap.add_argument("output", type=Path)
    args = ap.parse_args()
    book = load_workbook(args.source, read_only=True, data_only=True)
    rows = []
    counts = {}
    for name in ("SD1", "SD3.Prioritized 240", "SD3.Lowest 240", "SD3.Highest 240"):
        sheet = book[name]
        header = list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
        growth_col = "Mean" if name == "SD1" else "Mean_Growth"
        count = 0
        for excel_row, values in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
            item = dict(zip(header, values))
            if item.get("SMILES") is None:
                continue
            growth = item[growth_col]
            if not isinstance(growth, (int, float)):
                raise ValueError((name, excel_row, "Missing experimental growth"))
            rows.append(dict(record_id=f"abaucin:{name}:{excel_row}", smiles=item["SMILES"],
                source_file=args.source.name, source_sheet=name, excel_row=excel_row,
                sample_name=item.get("Name", item.get("Sample")), mean_growth=growth,
                label=int(growth < 0.2), label_rule="Mean_Growth < 0.2; >80% inhibition at 50 uM",
                strain="ATCC 17978", medium="LB", concentration_uM=50,
                study="doi:10.1038/s41589-023-01349-8",
                selection="primary_library" if name == "SD1" else "previous_model_selected"))
            count += 1
        counts[name] = count
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("x", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    print(json.dumps(counts, indent=2))


if __name__ == "__main__":
    main()
